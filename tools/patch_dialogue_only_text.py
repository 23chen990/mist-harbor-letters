#!/usr/bin/env python3
"""Move unquoted stage directions out of NPC台词. Player-facing text stays dialogue-only."""

from __future__ import annotations

import re
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "outputs/01a037d8-c4f0-70d3-bc69-8c3819094801/雾港来信_剧情作者工作簿.xlsx"
SHEET = "剧情剧本"
KEEP_AS_OBJECT_TEXT = {"线索查看"}
SKIP_TYPES = {
    "玩家选项",
    "位置选择",
    "稿件选择",
    "一次性观察",
    "场景探索",
    "场景行动",
    "纯跳转",
    "纯画面",
    "结局",
    "写稿",
    "判断选项",
    "限时观察",
}
SPEAKERS = (
    "版面编辑",
    "编辑",
    "沈砚舟",
    "林怀安",
    "林玉棠",
    "玉棠",
    "方仲山",
    "陈九生",
    "许济川",
    "阿成",
    "顾承钧",
    "场务",
    "伙计",
    "警员",
    "有人",
    "另一人",
)


def headers(ws) -> dict[str, int]:
    return {cell.value: i for i, cell in enumerate(ws[1]) if cell.value}


def cell(ws, row: int, col: dict[str, int], field: str):
    return ws.cell(row, col[field] + 1)


def split_lines(text: str) -> list[str]:
    normalized = text.replace("\r\n", "\n").replace("\r", "\n").replace("\\n", "\n")
    return [line.strip() for line in normalized.split("\n") if line.strip()]


def split_cue_and_quote(line: str) -> tuple[str, str] | None:
    match = re.search(r"[「『]", line)
    if not match:
        return None
    cue = line[: match.start()].strip()
    quote = line[match.start() :].strip()
    if not quote:
        return None
    cue = cue.rstrip("：:").strip()
    if not cue:
        return "", quote
    for speaker in sorted(SPEAKERS, key=len, reverse=True):
        if cue == speaker or cue == f"版面{speaker}":
            return "", f"{speaker}：{quote}" if not quote.startswith(speaker) else f"{cue}：{quote}"
        # 动作提示里出现两个人时，用最后点名的人当说话者（「玉棠想扶，许济川挡开」）。
        last_hit = -1
        last_speaker = ""
        for named in SPEAKERS:
            position = cue.rfind(named)
            if position > last_hit:
                last_hit = position
                last_speaker = named
        if last_speaker:
            spoken_as = "林玉棠" if last_speaker == "玉棠" else last_speaker
            return cue, f"{spoken_as}：{quote}"
    return cue, quote


def split_dialogue_and_visual(source: str) -> tuple[str, list[str]]:
    dialogue: list[str] = []
    visual: list[str] = []
    for line in split_lines(source):
        parsed = split_cue_and_quote(line)
        if parsed is None:
            visual.append(line)
            continue
        cue, spoken = parsed
        if cue:
            visual.append(cue)
        if spoken:
            dialogue.append(spoken)
    return "\n".join(dialogue), visual


def merge_visual(existing: str, extras: list[str]) -> str:
    parts = split_lines(existing) if existing else []
    for item in extras:
        if item and item not in parts:
            parts.append(item)
    return "\n".join(parts)


def main() -> int:
    wb = openpyxl.load_workbook(WORKBOOK)
    ws = wb[SHEET]
    col = headers(ws)
    changed = 0
    for row in range(2, ws.max_row + 1):
        kind = str(cell(ws, row, col, "内容类型").value or "")
        if kind in KEEP_AS_OBJECT_TEXT or kind in SKIP_TYPES:
            continue
        source = str(cell(ws, row, col, "NPC台词").value or "")
        if not source.strip():
            continue
        dialogue, visual = split_dialogue_and_visual(source)
        if dialogue == source.replace("\\n", "\n").strip() and not visual:
            continue
        if dialogue != source.replace("\\n", "\n").strip():
            cell(ws, row, col, "NPC台词").value = dialogue
            existing = str(cell(ws, row, col, "画面表现").value or "")
            cell(ws, row, col, "画面表现").value = merge_visual(existing, visual)
            changed += 1
            row_id = cell(ws, row, col, "自动ID").value
            print(f"{row_id}: dialogue-only, moved {len(visual)} visual notes")
    wb.save(WORKBOOK)
    print(f"PATCHED {changed} rows")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
