#!/usr/bin/env python3
"""Patch v6 workbook for rope-branch prototype (DEC-20260922-BRANCHING-UNFREEZE)."""

from __future__ import annotations

import sys
from copy import copy
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "outputs/01a037d8-c4f0-70d3-bc69-8c3819094801/雾港来信_剧情作者工作簿.xlsx"
SHEET = "剧情剧本"


def _headers(ws) -> dict[str, int]:
    return {cell.value: index for index, cell in enumerate(ws[1]) if cell.value}


def _set(ws, row: int, col: dict[str, int], field: str, value) -> None:
    ws.cell(row, col[field] + 1, value)


def _find_row(ws, col: dict[str, int], row_id: str) -> int:
    for row in range(2, ws.max_row + 1):
        if ws.cell(row, col["自动ID"] + 1).value == row_id:
            return row
    raise KeyError(row_id)


def _insert_after(ws, after_row: int) -> int:
    ws.insert_rows(after_row + 1)
    for col_idx in range(1, ws.max_column + 1):
        source = ws.cell(after_row, col_idx)
        target = ws.cell(after_row + 1, col_idx)
        target._style = copy(source._style)
    return after_row + 1


def patch_workbook(path: Path) -> None:
    wb = openpyxl.load_workbook(path)
    ws = wb[SHEET]
    col = _headers(ws)

    c21_row = _find_row(ws, col, "C21_004")
    _set(ws, c21_row, col, "自动ID", "C21_004A")
    _set(ws, c21_row, col, "出现条件", "report_focus=rope and police_told_rope=true")
    _set(
        ws,
        c21_row,
        col,
        "画面表现",
        "春和天桥吊绳状态存疑\n铁梯与绳排当晚封存。警方记录了整排收束状态。",
    )
    _set(ws, c21_row, col, "动作/表情备注", "春和对外只说仍在配合警方。")
    _set(
        ws,
        c21_row,
        col,
        "状态写入（不显示）",
        "published_focus=rope; evidence.rope_verifiable=true; chapter1_aftermath.rope=secured",
    )

    c21b_row = _insert_after(ws, c21_row)
    _set(ws, c21b_row, col, "自动ID", "C21_004B")
    _set(ws, c21b_row, col, "场景ID", "C21_NEWSPAPER")
    _set(ws, c21b_row, col, "场景", "次日清晨·报纸版面")
    _set(ws, c21b_row, col, "对话阶段", "第一章·次日见报")
    _set(ws, c21b_row, col, "阶段顺序", 1)
    _set(ws, c21b_row, col, "人物", "报纸")
    _set(ws, c21b_row, col, "内容类型", "纯画面")
    _set(ws, c21b_row, col, "话题", "实际见报")
    _set(ws, c21b_row, col, "出现条件", "report_focus=rope and police_told_rope=false")
    _set(ws, c21b_row, col, "下一话题", "第一章·第一章结尾")
    _set(
        ws,
        c21b_row,
        col,
        "画面表现",
        "春和天桥吊绳状态存疑\n次日再去春和，那根绳的收束位置已经改变，现场无法按昨晚所见原样核对。",
    )
    _set(ws, c21b_row, col, "动作/表情备注", "戏院场务不再允许报馆的人单独接近绳排。")
    _set(ws, c21b_row, col, "设计备注", "只显示玩家实际写出的版面结果。")
    _set(
        ws,
        c21b_row,
        col,
        "状态写入（不显示）",
        "published_focus=rope; evidence.rope_verifiable=false; chapter1_aftermath.rope=lost",
    )
    _set(ws, c21b_row, col, "是否说出口", "否")
    _set(ws, c21b_row, col, "演出/交互方式", "报纸UI；背景UI无对话框")

    c22_row = _find_row(ws, col, "C22_001")
    c22_secured = _insert_after(ws, c22_row)
    _set(ws, c22_secured, col, "自动ID", "C22_002")
    _set(ws, c22_secured, col, "场景ID", "C22_END")
    _set(ws, c22_secured, col, "场景", "《雾港日报》·赵敬文旧桌")
    _set(ws, c22_secured, col, "对话阶段", "第一章·第一章结尾")
    _set(ws, c22_secured, col, "阶段顺序", 2)
    _set(ws, c22_secured, col, "人物", "系统")
    _set(ws, c22_secured, col, "内容类型", "纯画面")
    _set(ws, c22_secured, col, "话题", "见报之后")
    _set(ws, c22_secured, col, "出现条件", "chapter1_aftermath.rope=secured")
    _set(
        ws,
        c22_secured,
        col,
        "画面表现",
        "这一篇见报之后：顾承钧依先前记录继续封存绳排；春和对外只说仍在配合警方。",
    )
    _set(ws, c22_secured, col, "设计备注", "DEC-20260922-BRANCHING-UNFREEZE 章末结账样板。")
    _set(ws, c22_secured, col, "是否说出口", "否")
    _set(ws, c22_secured, col, "演出/交互方式", "背景UI无对话框")

    c22_lost = _insert_after(ws, c22_secured)
    _set(ws, c22_lost, col, "自动ID", "C22_003")
    _set(ws, c22_lost, col, "场景ID", "C22_END")
    _set(ws, c22_lost, col, "场景", "《雾港日报》·赵敬文旧桌")
    _set(ws, c22_lost, col, "对话阶段", "第一章·第一章结尾")
    _set(ws, c22_lost, col, "阶段顺序", 2)
    _set(ws, c22_lost, col, "人物", "系统")
    _set(ws, c22_lost, col, "内容类型", "纯画面")
    _set(ws, c22_lost, col, "话题", "见报之后")
    _set(ws, c22_lost, col, "出现条件", "chapter1_aftermath.rope=lost")
    _set(
        ws,
        c22_lost,
        col,
        "画面表现",
        "这一篇见报之后：春和天桥昨日那排绳已不再保持昨晚的收束；报馆的人进后台会被场务拦在绳排之外。",
    )
    _set(ws, c22_lost, col, "设计备注", "DEC-20260922-BRANCHING-UNFREEZE 章末结账样板。")
    _set(ws, c22_lost, col, "是否说出口", "否")
    _set(ws, c22_lost, col, "演出/交互方式", "背景UI无对话框")

    wb.save(path)
    print(f"PATCHED: {path} ({ws.max_row - 1} story rows)")


def main() -> int:
    if not WORKBOOK.exists():
        print(f"Missing workbook: {WORKBOOK}", file=sys.stderr)
        return 1
    patch_workbook(WORKBOOK)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
